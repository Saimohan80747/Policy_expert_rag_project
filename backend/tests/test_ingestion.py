import io
import zipfile

import pymupdf as fitz
import pytest
from docx import Document
from fastapi import UploadFile

import embeddings
import zip_file_breaker as zfb
from ingestion import ingest_file
from retrieve_docs import hybrid_scale, retrieve_documents
from supabase_.delete_documents import delete_document

NS = "docs"


def upload(name, data):
    if isinstance(data, str):
        data = data.encode("utf-8")
    return UploadFile(filename=name, file=io.BytesIO(data))


def long_text(tag, paragraphs=12):
    return "\n\n".join(f"{tag} paragraph {i}. " + ("policy detail " * 40) for i in range(paragraphs))


def texts(env):
    return [r["metadata"]["text"] for r in env.index.ns.get(NS, {}).values()]


def test_new_document_is_ingested(env):
    result = ingest_file(upload("Leave Policy.txt", "Casual leave is 12 days per year."))
    assert result["status"] == "ingested"
    assert result["file_name"] == "Leave_Policy.txt"
    assert result["version"] == 1 and result["chunks"] == 1
    assert env.index.all_ids() == [f"{result['doc_id']}:v1:0"]
    assert "Leave_Policy.txt" in env.supabase.files
    row = env.supabase.tables["documents"][0]
    assert row["is_current"] is True and row["chunk_count"] == 1 and row["doc_type"] == "txt"


def test_same_bytes_is_duplicate(env):
    ingest_file(upload("a.txt", "same content"))
    again = ingest_file(upload("renamed.txt", "same content"))
    assert again["status"] == "duplicate"
    assert again["duplicate_of"] == "a.txt"
    assert len(env.supabase.tables["documents"]) == 1
    assert len(env.index.all_ids()) == 1


def test_changed_content_becomes_new_version(env):
    v1 = ingest_file(upload("policy.txt", "Notice period is 30 days."))
    v2 = ingest_file(upload("policy.txt", "Notice period is 60 days."))
    assert (v1["version"], v2["version"], v2["supersedes"]) == (1, 2, 1)
    assert v1["doc_id"] == v2["doc_id"]

    meta = {i: r["metadata"] for i, r in env.index.ns[NS].items()}
    assert meta[f"{v1['doc_id']}:v1:0"]["is_current"] is False
    assert meta[f"{v1['doc_id']}:v2:0"]["is_current"] is True
    rows = {r["version"]: r["is_current"] for r in env.supabase.tables["documents"]}
    assert rows == {1: False, 2: True}


def test_search_only_returns_current_version(env):
    ingest_file(upload("policy.txt", "Notice period is 30 days."))
    ingest_file(upload("policy.txt", "Notice period is 60 days."))
    current = retrieve_documents("notice period")
    assert [r["version"] for r in current] == [2]
    assert {r["version"] for r in retrieve_documents("notice period", current_only=False)} == {1, 2}


def test_older_versions_are_retired_with_per_vector_fallback(env):
    env.index.filter_update_supported = False  # e.g. an older SDK
    v1 = ingest_file(upload("policy.txt", "Notice period is 30 days."))
    ingest_file(upload("policy.txt", "Notice period is 60 days."))
    meta = {i: r["metadata"]["is_current"] for i, r in env.index.ns[NS].items()}
    assert meta == {f"{v1['doc_id']}:v1:0": False, f"{v1['doc_id']}:v2:0": True}


def test_search_drops_versions_pinecone_has_not_retired_yet(env):
    ingest_file(upload("policy.txt", "Notice period is 30 days."))
    ingest_file(upload("policy.txt", "Notice period is 60 days."))
    for rec in env.index.ns[NS].values():  # simulate Pinecone's async update lag
        rec["metadata"]["is_current"] = True
    assert [r["version"] for r in retrieve_documents("notice period")] == [2]


def test_chunks_carry_the_public_file_url(env):
    ingest_file(upload("policy.txt", "Notice period is 30 days."))
    (rec,) = env.index.ns[NS].values()
    assert rec["metadata"]["file_url"] == "https://fake.storage/documents/policy.txt"


def test_unsupported_and_empty_files_fail(env):
    assert ingest_file(upload("tool.exe", b"MZ"))["status"] == "failed"
    empty = ingest_file(upload("empty.txt", b""))
    assert empty["status"] == "failed" and "empty" in empty["error"].lower()
    assert env.index.all_ids() == []
    assert env.supabase.tables.get("documents", []) == []


def test_csv_rows_become_records_and_cp1252_is_decoded(env):
    data = "Name,City\nJosé,Zürich\nAsha,Pune\n".encode("cp1252")
    assert ingest_file(upload("people.csv", data))["status"] == "ingested"
    (text,) = texts(env)
    assert "Name: José | City: Zürich" in text
    assert "Name: Asha | City: Pune" in text


def test_xlsx_sheets_are_extracted(env):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Budget"
    wb.active.append(["Dept", "Amount"])
    wb.active.append(["Platform", 1200])
    buf = io.BytesIO()
    wb.save(buf)
    assert ingest_file(upload("budget.xlsx", buf.getvalue()))["status"] == "ingested"
    (text,) = texts(env)
    assert "## Sheet: Budget" in text and "Dept: Platform | Amount: 1200" in text


def test_docx_keeps_headings_and_tables(env):
    doc = Document()
    doc.add_heading("Leave Policy", 1)
    doc.add_paragraph("Applies to all employees.")
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate([("Type", "Days"), ("Casual", "12")]):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    buf = io.BytesIO()
    doc.save(buf)
    assert ingest_file(upload("leave.docx", buf.getvalue()))["status"] == "ingested"
    (text,) = texts(env)
    assert "# Leave Policy" in text and "Casual | 12" in text


def test_pdf_text_and_image_only_pdf(env):
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Expense limit is 5000 rupees")
    assert ingest_file(upload("expense.pdf", pdf.tobytes()))["status"] == "ingested"
    assert "Expense limit is 5000 rupees" in texts(env)[0]

    blank = fitz.open()
    blank.new_page()
    result = ingest_file(upload("scan.pdf", blank.tobytes()))
    assert result["status"] == "failed" and "OCR" in result["error"]


def test_corrupt_pdf_fails_cleanly(env):
    result = ingest_file(upload("broken.pdf", b"%PDF-1.7 not really a pdf"))
    assert result["status"] == "failed" and result["step"] == "extract_text"


def test_long_document_is_split_into_deterministic_chunks(env):
    result = ingest_file(upload("handbook.txt", long_text("Handbook")))
    assert result["chunks"] > 3
    expected = [f"{result['doc_id']}:v1:{i}" for i in range(result["chunks"])]
    assert sorted(env.index.all_ids()) == sorted(expected)
    assert all(len(t) <= 2000 for t in texts(env))


def test_partial_upsert_failure_leaves_nothing_behind(env, monkeypatch):
    monkeypatch.setattr(embeddings, "UPSERT_BATCH", 2)
    env.index.fail_on_upsert_call = 2
    result = ingest_file(upload("handbook.txt", long_text("Handbook")))
    assert result["status"] == "failed" and result["step"] == "embedding_chunks"
    assert env.index.all_ids() == []
    assert env.supabase.tables.get("documents", []) == []


def test_cleanup_failure_does_not_hide_the_original_error(env, monkeypatch):
    monkeypatch.setattr(embeddings, "UPSERT_BATCH", 2)
    env.index.fail_on_upsert_call = 1
    env.index.fail_list = True  # cleanup itself fails too
    result = ingest_file(upload("handbook.txt", long_text("Handbook")))
    assert result["status"] == "failed" and "pinecone down" in result["error"]


def test_wrong_index_type_gives_a_clear_error(monkeypatch):
    from types import SimpleNamespace

    fake_pc = SimpleNamespace(
        has_index=lambda name: True,
        describe_index=lambda name: SimpleNamespace(dimension=None, metric=None),
    )
    monkeypatch.setattr(embeddings, "get_pc", lambda: fake_pc)
    monkeypatch.setattr(embeddings, "_index", None)
    with pytest.raises(RuntimeError, match="document schema"):
        embeddings.get_index()


def test_metadata_failure_removes_vectors(env):
    env.supabase.fail_upload = True
    result = ingest_file(upload("a.txt", "some text"))
    assert result["status"] == "failed" and result["step"] == "store_document"
    assert env.index.all_ids() == []
    assert env.supabase.tables.get("documents", []) == []


def test_failed_ingest_can_be_retried(env):
    env.supabase.fail_upload = True
    assert ingest_file(upload("a.txt", "some text"))["status"] == "failed"
    env.supabase.fail_upload = False
    assert ingest_file(upload("a.txt", "some text"))["status"] == "ingested"


@pytest.mark.parametrize("legacy_list", [False, True])
def test_listing_works_with_old_and_new_sdk_page_shapes(env, legacy_list):
    env.index.legacy_list = legacy_list
    result = ingest_file(upload("handbook.txt", long_text("Handbook")))
    ids = embeddings.list_ids(result["doc_id"])
    assert len(ids) == result["chunks"] and all(isinstance(i, str) for i in ids)
    assert embeddings.delete_doc_vectors(result["doc_id"]) == result["chunks"]
    assert env.index.all_ids() == []


def test_delete_document_removes_everything(env):
    ingest_file(upload("policy.txt", "Notice period is 30 days."))
    ingest_file(upload("policy.txt", "Notice period is 60 days."))
    ingest_file(upload("other.txt", "Unrelated document."))
    result = delete_document("policy.txt")
    assert result["status"] == "deleted"
    assert result["versions_deleted"] == 2 and result["vectors_deleted"] == 2
    assert "policy.txt" not in env.supabase.files
    assert len(env.index.all_ids()) == 1
    assert [r["file_name"] for r in env.supabase.tables["documents"]] == ["other.txt"]
    assert delete_document("policy.txt")["status"] == "not_found"


def make_zip(entries):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, data in entries.items():
            z.writestr(name, data)
    return buf.getvalue()


def test_zip_with_nested_zip_duplicates_junk_and_bad_files(env):
    nested = make_zip({"b.txt": "nested file content"})
    archive = make_zip(
        {
            "docs/a.txt": "first file content",
            "docs/a_copy.txt": "first file content",
            "inner/nested.zip": nested,
            "__MACOSX/._a.txt": "junk",
            ".DS_Store": "junk",
            "tool.exe": "MZ",
        }
    )
    result = zfb.zip_file_breaker(upload("bundle.zip", archive), ingest_file)
    assert sorted(e["result"]["file_name"] for e in result["ingested_files"]) == ["a.txt", "b.txt"]
    assert [e["filename"] for e in result["duplicate_files"]] == ["docs/a_copy.txt"]
    assert [e["filename"] for e in result["failed_files"]] == ["tool.exe"]
    assert len(env.index.all_ids()) == 2


class _NoSeekable:
    """File-like object without `seekable`, like SpooledTemporaryFile on Python < 3.11."""

    def __init__(self, data):
        self._f = io.BytesIO(data)
        self.read, self.seek, self.tell = self._f.read, self._f.seek, self._f.tell


def test_zip_upload_whose_file_object_lacks_seekable(env):
    archive = make_zip({"x.txt": "first file", "y.bin": "nope"})
    result = zfb.zip_file_breaker(
        UploadFile(filename="pack.zip", file=_NoSeekable(archive)), ingest_file
    )
    assert [e["result"]["file_name"] for e in result["ingested_files"]] == ["x.txt"]
    assert len(result["failed_files"]) == 1


def test_zip_size_limit_and_invalid_zip(env, monkeypatch):
    monkeypatch.setattr(zfb, "MAX_ZIP_UNCOMPRESSED_BYTES", 10)
    big = zfb.zip_file_breaker(upload("big.zip", make_zip({"a.txt": "x" * 100})), ingest_file)
    assert "limit" in big["failed_files"][0]["error"]
    bad = zfb.zip_file_breaker(upload("bad.zip", b"not a zip"), ingest_file)
    assert bad["failed_files"][0]["error"].startswith("Not a valid ZIP file")


def test_hybrid_scale():
    dense, sparse = hybrid_scale([1.0, 2.0], {"indices": [3], "values": [4.0]}, 0.75)
    assert dense == [0.75, 1.5] and sparse == {"indices": [3], "values": [1.0]}
    with pytest.raises(ValueError):
        hybrid_scale([1.0], {"indices": [], "values": []}, 1.5)
